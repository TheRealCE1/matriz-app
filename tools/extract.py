"""Genera server/data/employees.json a partir del Excel de matrices."""
import sys, re, json, pathlib
import openpyxl

src = sys.argv[1] if len(sys.argv) > 1 else 'MATRICES_EN_INTRANET_CC_1.xlsm'
out = pathlib.Path(__file__).resolve().parent.parent / 'server' / 'data' / 'employees.json'
wb = openpyxl.load_workbook(src, data_only=True, keep_vba=True)

SKIP = {'Menú','Documentación','Resumen','.','..','AVANCE MENSUAL','Hoja47','Base graficas','2X2 CDCE','Soldadura básica'}
LINEA = {'línea en tress','linea en tress','línea','linea','línea tress','linea tress'}

def find_header(ws):
    for r in range(1, min(ws.max_row, 20) + 1):
        for c in range(1, min(ws.max_column, 10) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().lower() in ('nómina', 'nomina'):
                return r
    return None

emps = {}
for name in wb.sheetnames:
    if name in SKIP: continue
    ws = wb[name]; hr = find_header(ws)
    if hr is None: continue
    h = {c: str(ws.cell(hr, c).value).strip() for c in range(1, ws.max_column + 1)
         if isinstance(ws.cell(hr, c).value, str) and ws.cell(hr, c).value.strip()}
    col = lambda f: next((c for c, v in h.items() if f(v.lower())), None)
    idc, nc = col(lambda v: v in ('nómina', 'nomina')), col(lambda v: v == 'nombre')
    tc, lc, pc = col(lambda v: v == 'turno'), col(lambda v: v in LINEA), col(lambda v: v == 'puesto')
    if not idc or not nc: continue
    meta = {idc, nc, tc, lc, pc}
    skills = [c for c in h if c not in meta and not re.fullmatch(r'\d+', h[c])]
    empty = 0; r = hr + 1
    while r <= ws.max_row and empty < 5:
        i, n = ws.cell(r, idc).value, ws.cell(r, nc).value
        if i is None and n is None: empty += 1; r += 1; continue
        empty = 0
        if n and str(i) != '#N/A':
            e = emps.setdefault(str(i), {'n': n, 'm': {}, 'c': []})
            sk = {h[c]: ws.cell(r, c).value for c in skills
                  if isinstance(ws.cell(r, c).value, (int, float)) and ws.cell(r, c).value}
            if sk:
                e['m'][name] = {'turno': ws.cell(r, tc).value if tc else None,
                                'linea': ws.cell(r, lc).value if lc else None,
                                'puesto': ws.cell(r, pc).value if pc else None, 'skills': sk}
        r += 1

ws = wb['Soldadura básica']
for r in range(2, ws.max_row + 1):
    i, n, f, c = (ws.cell(r, k).value for k in range(1, 5))
    if i is None or not c: continue
    e = emps.setdefault(str(i), {'n': n or f'Empleado #{i}', 'm': {}, 'c': []})
    e['c'].append({'curso': str(c).strip(' "'), 'fecha': f.strftime('%Y-%m-%d') if hasattr(f, 'strftime') else None})

for k, e in emps.items():
    if not isinstance(e['n'], str) or e['n'].startswith('#'):
        e['n'] = f'Empleado #{k}'

out.write_text(json.dumps(emps, ensure_ascii=False), encoding='utf-8')
print(f'{len(emps)} colaboradores -> {out}')
